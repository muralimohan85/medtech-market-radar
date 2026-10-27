"""Export a published version to Word, PowerPoint and Excel.

audience="internal" keeps everything your licences allow inside your organisation.
audience="external" drops sentences that rest on licensed or uploaded material and
puts a placeholder in their place, so a deck can go to a partner without breaching a
data licence.
"""

from __future__ import annotations

import io
from datetime import date

from radar.config import SECTION_ORDER
from radar.report.render import Citations
from radar.safety import licence
from radar.schemas import Report, Sentence
from radar.store import Store

ACCENT = (0x0B, 0x61, 0x76)


def allowed(s: Sentence, store: Store, audience: str) -> bool:
    return all(licence.can_export(f.licence, audience) for f in (store.fact(c) for c in s.cites) if f)


def filtered(rep: Report, store: Store, audience: str) -> tuple[Report, int]:
    """A copy of the report with non-exportable sentences replaced by a placeholder."""
    out = Report.from_dict(rep.to_dict())
    removed = 0
    for sec in out.sections.values():
        for b in sec.blocks:
            keep = []
            for s in b.sentences:
                if allowed(s, store, audience):
                    keep.append(s)
                else:
                    removed += 1
                    if not keep or keep[-1].text != licence.RESTRICTED:
                        keep.append(Sentence(text=licence.RESTRICTED, cites=[], kind="restricted"))
            b.sentences = keep
    return out, removed


def _sentence_text(s: Sentence, cites: Citations) -> str:
    marks = cites.marks(s)
    return s.text + ("".join(f"[{n}]" for n in marks) if marks else "")


def to_docx(rep: Report, store: Store, title: str, audience: str = "internal", version: int | None = None) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    rep, removed = filtered(rep, store, audience)
    cites = Citations(rep, store, SECTION_ORDER)
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    h = d.add_heading(title, 0)
    h.runs[0].font.color.rgb = RGBColor(*ACCENT)
    d.add_paragraph(f"Version {version or ''} · as of {rep.as_of} · {audience} distribution"
                    + (f" · {removed} licensed statements withheld" if removed else ""))
    for sid in SECTION_ORDER:
        sec = rep.sections.get(sid)
        if not sec:
            continue
        d.add_heading(sec.title, 1)
        for b in sec.blocks:
            if b.type == "h3":
                d.add_heading(b.text, 2)
            elif b.sentences:
                p = d.add_paragraph()
                for s in b.sentences:
                    r = p.add_run(_sentence_text(s, cites) + " ")
                    if s.kind == "judgement":
                        r.italic = True
    d.add_heading("8 Sources", 1)
    for src in cites.sources():
        d.add_paragraph(f"[{src['n']}] {src['title']} · {src['licence_label']} · {src.get('published', '')}"
                        + (f" · {src['url']}" if src.get("url") and src["kind"] == "public" else ""))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_pptx(rep: Report, store: Store, title: str, events: list, audience: str = "internal") -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    rep, _ = filtered(rep, store, audience)
    cites = Citations(rep, store, SECTION_ORDER)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def slide(head: str, lines: list[str]):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = head
        s.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(*ACCENT)
        tb = s.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(11.9), Inches(5.4)).text_frame
        tb.word_wrap = True
        for i, ln in enumerate(lines[:9]):
            p = tb.paragraphs[0] if i == 0 else tb.add_paragraph()
            p.text = ln
            p.font.size = Pt(16)
            p.space_after = Pt(8)

    t = prs.slides.add_slide(prs.slide_layouts[0])
    t.shapes.title.text = title
    t.placeholders[1].text = f"As of {rep.as_of} · {audience} distribution"
    for sid in ("summary", "market", "competition", "pricing", "pipeline"):
        sec = rep.sections.get(sid)
        if sec:
            lines = []
            for b in sec.blocks:
                lines += [f"▸ {b.text}"] if b.type == "h3" else [_sentence_text(s, cites) for s in b.sentences]
            slide(sec.title, lines)
    if events:
        slide("What changed", [f"[{e.level.upper()}] {e.date} {e.title}" for e in events])
    slide("Sources", [f"[{s['n']}] {s['title'][:90]} ({s['licence_label']})" for s in cites.sources()])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def to_xlsx(rep: Report, store: Store, events: list, audience: str = "internal") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="0B6176"), Font(color="FFFFFF", bold=True)

    def sheet(name: str, header: list[str], rows: list[list]):
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(70, max(10, *(len(str(c.value or "")) for c in col)) + 2)

    wb.remove(wb.active)
    ok = lambda f: licence.can_export(f.licence, audience)  # noqa: E731
    est = [f for f in store.facts(as_of=rep.as_of, topic=["market_estimate", "upload"]) if f.data.get("metric") == "market_value"]
    sheet("Market estimates", ["Source", "Licence", "Date", "Region", "Year", "Low (bn)", "High (bn)", "CAGR %", "Period"],
          [[(store.source(f.source_id) or {}).get("title", ""), licence.label(f.licence), f.as_of, f.data.get("region"),
            f.data.get("year"), f.data.get("value_low"), f.data.get("value_high"), f.data.get("cagr"), f.data.get("cagr_period")]
           if ok(f) else ["[restricted]", licence.label(f.licence), f.as_of] for f in est])
    sheet("What changed", ["Detected", "Published", "Level", "Kind", "Company", "Title", "Sections", "Why"],
          [[e.detected, e.date, e.level, e.kind, e.entity, e.title, ", ".join(e.sections), "; ".join(e.reasons)] for e in events])
    prices = [f for f in store.facts(as_of=rep.as_of) if "price_low" in f.data]
    sheet("Prices", ["Date", "Company", "Low", "High", "Currency", "Statement", "Licence"],
          [[f.as_of, f.entity, f.data["price_low"], f.data["price_high"], f.data.get("currency"),
            f.claim if ok(f) else "[restricted]", licence.label(f.licence)] for f in prices])
    sheet("Sources", ["Id", "Connector", "Title", "Published", "Licence", "URL"],
          [[s["id"], s["connector"], s["title"], s["published"], licence.label(s["licence"]),
            s["url"] if licence.kind_of(s["licence"]) == "public" else ""] for s in store.sources()])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def stamp() -> str:
    return date.today().isoformat()
