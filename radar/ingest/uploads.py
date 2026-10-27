"""Read uploaded reports (PDF, PowerPoint, Excel, Word, CSV, text) into passages and candidate facts.

Uploads are the analyst's own material: purchased reports, tender summaries, price
lists. They are read locally, tagged with the licence the analyst chose, and never
leave the workspace in external exports. The deterministic extractor picks out
market-size, growth and price statements; the Claude engine can extract more.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from radar.safety.sanitizer import clean, sentences

SUPPORTED = {".pdf", ".pptx", ".xlsx", ".xlsm", ".docx", ".csv", ".txt", ".md"}


@dataclass
class Passage:
    where: str           # "page 2", "slide 4", "sheet Prices row 7"
    text: str


@dataclass
class Parsed:
    name: str
    passages: list[Passage] = field(default_factory=list)
    quarantined: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n".join(p.text for p in self.passages)


def parse(path: str | Path) -> Parsed:
    p = Path(path)
    ext = p.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"unsupported file type {ext}; use one of {', '.join(sorted(SUPPORTED))}")
    raw: list[Passage] = []
    if ext == ".pdf":
        import pdfplumber

        with pdfplumber.open(p) as pdf:
            raw = [Passage(f"page {i}", pg.extract_text() or "") for i, pg in enumerate(pdf.pages, 1)]
    elif ext == ".pptx":
        from pptx import Presentation

        for i, slide in enumerate(Presentation(p).slides, 1):
            txt = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.has_text_frame]
            raw.append(Passage(f"slide {i}", "\n".join(txt)))
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for r, row in enumerate(ws.iter_rows(values_only=True), 1):
                cells = [str(c) for c in row if c not in (None, "")]
                if cells:
                    raw.append(Passage(f"sheet {ws.title} row {r}", " | ".join(cells)))
    elif ext == ".docx":
        import docx

        raw = [Passage(f"paragraph {i}", para.text) for i, para in enumerate(docx.Document(p).paragraphs, 1) if para.text.strip()]
    elif ext == ".csv":
        with open(p, newline="", encoding="utf-8", errors="replace") as fh:
            raw = [Passage(f"row {i}", " | ".join(row)) for i, row in enumerate(csv.reader(fh), 1) if any(row)]
    else:
        raw = [Passage("text", p.read_text(encoding="utf-8", errors="replace"))]

    out = Parsed(name=p.name)
    for ps in raw:
        c = clean(_join_lines(ps.text))
        out.quarantined += c.quarantined
        if c.text.strip():
            out.passages.append(Passage(ps.where, c.text))
    return out


def _join_lines(text: str) -> str:
    """Join wrapped lines into prose; short lines with no end punctuation are headings and become their own sentence."""
    lines = [ln.strip() for ln in (text or "").splitlines() if ln.strip()]
    out: list[str] = []
    for i, ln in enumerate(lines):
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        heading = len(ln) < 60 and not re.search(r"[.!?:;,]$", ln) and (not nxt or nxt[:1].isupper())
        out.append(ln + ("." if heading and not ln.endswith(".") else ""))
    return " ".join(out)


# ------------------------------------------------------------------ fact extraction
_NUM = r"(\d+(?:\.\d+)?)"
_CUR = r"(USD|US\$|\$|EUR|€|INR|₹|GBP|£)"
MARKET = re.compile(rf"{_CUR}\s?{_NUM}(?:\s?(?:-|–|to)\s?(?:{_CUR}\s?)?{_NUM})?\s?(billion|bn|million|mn|m)\b", re.I)
CAGR = re.compile(rf"(?:CAGR of\s?{_NUM}\s?%|{_NUM}\s?%\s?(?:CAGR|compound annual|a year|per year|annually))", re.I)
PRICE = re.compile(rf"{_CUR}\s?{_NUM}(?:\s?(?:-|–|to)\s?(?:{_CUR}\s?)?{_NUM})?\s?per test", re.I)
YEAR = re.compile(r"\b(20[2-4]\d)\b")
PERIOD = re.compile(r"\b(20[2-4]\d)\s?(?:-|–|to)\s?(20[2-4]\d)\b")


def _cur(s: str) -> str:
    return {"$": "USD", "US$": "USD", "€": "EUR", "₹": "INR", "£": "GBP"}.get(s.upper() if s else "", (s or "USD").upper())


def extract_facts(parsed: Parsed) -> list[dict]:
    """Candidate facts: {topic, claim, where, data}. Every claim is a sentence from the upload."""
    facts = []
    for ps in parsed.passages:
        for s in sentences(ps.text):
            low = s.lower()
            m = MARKET.search(s)
            if m and "market" in low and "per test" not in low:
                scale = 1.0 if m.group(5).lower() in ("billion", "bn") else 0.001
                lo = float(m.group(2)) * scale
                hi = float(m.group(4)) * scale if m.group(4) else lo
                years = YEAR.findall(s)
                period = PERIOD.search(s)
                c = CAGR.search(s)
                data = {"metric": "market_value", "value_low": round(lo, 3), "value_high": round(hi, 3),
                        "unit": f"{_cur(m.group(1))} bn", "year": int(years[0]) if years else None, "region": "Global"}
                if c:
                    data["cagr"] = float(c.group(1) or c.group(2))
                if period:
                    data["cagr_period"] = f"{period.group(1)}-{period.group(2)}"
                facts.append({"topic": "market_estimate", "claim": s, "where": ps.where, "data": data})
                continue
            p = PRICE.search(s)
            if p:
                lo = float(p.group(2))
                hi = float(p.group(4)) if p.group(4) else lo
                region = "United States" if re.search(r"\b(US|U\.S\.|United States)\b", s) else (
                    "Europe" if re.search(r"Europe|EU\b", s) else "")
                facts.append({"topic": "pricing", "claim": s, "where": ps.where,
                              "data": {"price_low": lo, "price_high": hi, "currency": _cur(p.group(1)), "region": region}})
    return facts
